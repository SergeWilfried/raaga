"""Convert the Nordgold stock sheet into seed JSON for seed-nordgold.ts.

Usage: python3 scripts/xlsx-to-nordgold-json.py <country> <sheet-location> <warehouse-location> <out-prefix> [block-size]

<sheet-location> is the value in the sheet's Location column; <warehouse-location> is
the name used for the warehouse ("Nordgold <warehouse-location>"), e.g. Lero -> Conakry.
Writes <out-prefix>.json (all rows) or, with a block size, <out-prefix>-01.json,
-02.json ... each holding at most that many rows. Rows are de-duplicated by
G-code (the product handle is built from it) and '#' placeholders become null.
"""
import json, os, sys, warnings
import openpyxl

warnings.filterwarnings("ignore")
country, sheet_location, location, prefix = sys.argv[1:5]
block = int(sys.argv[5]) if len(sys.argv) > 5 else 0

ws = openpyxl.load_workbook("seed/Nordgold_stock_2026_WA.xlsx", read_only=True)["WA DS 2026"]
clean = lambda v: None if v is None or str(v).strip() in ("", "#") else str(v).strip()

rows, seen, dupes = [], set(), 0
for r in ws.iter_rows(min_row=7, values_only=True):
    if r[0] != country or r[2] != sheet_location or not clean(r[6]):
        continue
    code = clean(r[6])
    if code in seen:
        dupes += 1
        continue
    seen.add(code)
    rows.append({
        "code": code,
        "description": clean(r[7]) or code,
        "brand": clean(r[3]),
        "equipment": clean(r[4]),
        "type": clean(r[5]),
        "mpn": clean(r[8]),
        "unit": clean(r[10]) or "Pieces",
        "country": country,
        "mine": clean(r[1]),
        "location": location,
        "qty": int(r[9] or 0),
    })

# Optional: SKIP=<n> drops the first n rows and FIRST_BLOCK=<k> numbers the
# first output file k, to re-split only the not-yet-loaded tail.
skip = int(os.environ.get("SKIP", 0))
first = int(os.environ.get("FIRST_BLOCK", 1))
rows = rows[skip:]

print(f"{len(rows)} rows, {dupes} duplicate G-codes skipped")
if not block:
    json.dump(rows, open(f"{prefix}.json", "w"), ensure_ascii=False)
else:
    n = (len(rows) + block - 1) // block
    for i in range(n):
        name = f"{prefix}-{first+i:02d}.json"
        json.dump(rows[i*block:(i+1)*block], open(name, "w"), ensure_ascii=False)
        print(name, len(rows[i*block:(i+1)*block]))
